"""Gera uma planilha de apoio visual (KPIs, tabelas-resumo e gráficos) a
partir do relatório tratado/cruzado, para conferência humana rápida.

Não substitui a planilha de dados completa: a aba "Dados" mantém TODAS as
colunas do relatório — os gráficos são um complemento, não um resumo que
descarta informação.
"""
from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.utils import converter_numero_br

logger = logging.getLogger("automacao_transportadora")

_FONTE_CABECALHO = Font(color="FFFFFF", bold=True)
_PREENCHIMENTO_CABECALHO = PatternFill("solid", fgColor="1F4E78")
_FONTE_TITULO = Font(bold=True, size=13)
_FONTE_KPI_ROTULO = Font(bold=True)
_FONTE_KPI_VALOR = Font(size=14, bold=True, color="1F4E78")
_PREENCHIMENTO_DUPLICADO = PatternFill("solid", fgColor="FFC7CE")

_SITUACOES_SEM_MOTORISTA = {"NÃO ENCONTRADO", "SEM MANIFESTO", "CONFLITO DE MANIFESTO"}


def _autoajustar_largura(ws, df: pd.DataFrame) -> None:
    for indice, coluna in enumerate(df.columns, start=1):
        maior = max(
            [len(str(coluna))] + [len(str(v)) for v in df[coluna].astype(str).head(300)]
        )
        ws.column_dimensions[get_column_letter(indice)].width = min(max(maior + 2, 10), 45)


def _escrever_dados(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Dados")
    ws.append(list(df.columns))
    for celula in ws[1]:
        celula.font = _FONTE_CABECALHO
        celula.fill = _PREENCHIMENTO_CABECALHO
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for linha in df.itertuples(index=False):
        ws.append(["" if pd.isna(v) else v for v in linha])

    ultima_linha = len(df) + 1
    ultima_coluna = len(df.columns)

    if ultima_linha > 1 and ultima_coluna > 0:
        referencia = f"A1:{get_column_letter(ultima_coluna)}{ultima_linha}"
        tabela = Table(displayName="TabelaDados", ref=referencia)
        tabela.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tabela)

    ws.freeze_panes = "A2"
    _autoajustar_largura(ws, df)

    if "dias_atraso" in df.columns and ultima_linha > 1:
        letra = get_column_letter(df.columns.get_loc("dias_atraso") + 1)
        ws.conditional_formatting.add(
            f"{letra}2:{letra}{ultima_linha}",
            ColorScaleRule(
                start_type="min", start_color="63BE7B",
                mid_type="percentile", mid_value=50, mid_color="FFEB84",
                end_type="max", end_color="F8696B",
            ),
        )

    if "duplicado" in df.columns and ultima_linha > 1:
        indice_coluna = df.columns.get_loc("duplicado") + 1
        for linha_numero in range(2, ultima_linha + 1):
            celula = ws.cell(row=linha_numero, column=indice_coluna)
            if str(celula.value).strip().lower() == "true":
                celula.fill = _PREENCHIMENTO_DUPLICADO


def _calcular_kpis(df: pd.DataFrame) -> list[tuple[str, object]]:
    kpis: list[tuple[str, object]] = [("Total de cargas", len(df))]

    if "motorista" in df.columns:
        sem_motorista = int(df["motorista"].isin(_SITUACOES_SEM_MOTORISTA).sum())
        kpis.append(("Cargas com motorista", len(df) - sem_motorista))
        kpis.append(("Cargas sem motorista/conflito", sem_motorista))

    if "peso_real" in df.columns:
        pesos = df["peso_real"].map(converter_numero_br).dropna()
        kpis.append(("Peso total (kg)", round(float(pesos.sum()), 2)))

    if "volumes" in df.columns:
        volumes = df["volumes"].map(converter_numero_br).dropna()
        kpis.append(("Total de volumes", int(volumes.sum())))

    if "valor_frete" in df.columns:
        fretes = df["valor_frete"].map(converter_numero_br).dropna()
        kpis.append(("Valor total de frete (R$)", round(float(fretes.sum()), 2)))

    if "dias_atraso" in df.columns:
        atraso = df["dias_atraso"].map(converter_numero_br)
        kpis.append(("Cargas com atraso (dias_atraso > 0)", int((atraso > 0).sum())))

    if "duplicado" in df.columns:
        duplicados = int(df["duplicado"].astype(str).str.lower().eq("true").sum())
        kpis.append(("Pedidos duplicados", duplicados))

    if "cliente" in df.columns:
        kpis.append(("Clientes distintos (destinatário)", int(df["cliente"].nunique())))

    if "remetente" in df.columns:
        kpis.append(("Remetentes distintos", int(df["remetente"].nunique())))

    if "motorista" in df.columns:
        motoristas_reais = df.loc[~df["motorista"].isin(_SITUACOES_SEM_MOTORISTA), "motorista"]
        kpis.append(("Motoristas distintos em rota", int(motoristas_reais.nunique())))

    return kpis


def _escrever_resumo(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Resumo", 0)
    ws["A1"] = "RESUMO DA ANÁLISE"
    ws["A1"].font = _FONTE_TITULO

    linha = 3
    for rotulo, valor in _calcular_kpis(df):
        ws.cell(row=linha, column=1, value=rotulo).font = _FONTE_KPI_ROTULO
        ws.cell(row=linha, column=2, value=valor).font = _FONTE_KPI_VALOR
        linha += 1

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 20


def _escrever_contagem_com_grafico(
    wb: Workbook,
    nome_aba: str,
    df: pd.DataFrame,
    coluna: str,
    titulo_grafico: str,
    limite: int = 15,
    tipo: str = "bar",
) -> None:
    if coluna not in df.columns:
        return

    contagem = (
        df[coluna].fillna("NÃO INFORMADO").astype(str).str.strip()
        .replace("", "NÃO INFORMADO")
        .value_counts()
        .head(limite)
    )
    if contagem.empty:
        return

    ws = wb.create_sheet(nome_aba)
    ws.append([coluna.replace("_", " ").capitalize(), "Quantidade"])
    for celula in ws[1]:
        celula.font = _FONTE_CABECALHO
        celula.fill = _PREENCHIMENTO_CABECALHO

    for nome, quantidade in contagem.items():
        ws.append([nome, int(quantidade)])

    ultima_linha = len(contagem) + 1
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 14

    dados_grafico = Reference(ws, min_col=2, min_row=1, max_row=ultima_linha)
    categorias = Reference(ws, min_col=1, min_row=2, max_row=ultima_linha)

    if tipo == "pizza":
        grafico = PieChart()
    else:
        grafico = BarChart()
        grafico.y_axis.title = "Quantidade"
        grafico.x_axis.title = coluna.replace("_", " ").capitalize()

    grafico.title = titulo_grafico
    grafico.add_data(dados_grafico, titles_from_data=True)
    grafico.set_categories(categorias)
    grafico.height = 10
    grafico.width = 20
    ws.add_chart(grafico, "D2")


def gerar_relatorio_visual(df: pd.DataFrame, caminho_saida: Path) -> Path:
    """Gera a planilha de análise visual (KPIs + tabelas-resumo + gráficos)
    a partir do relatório tratado/cruzado (`resultados['relatorio_completo']`
    do cruzamento, por exemplo) e salva em `caminho_saida`.
    """
    wb = Workbook()
    wb.remove(wb.active)

    _escrever_resumo(wb, df)
    _escrever_dados(wb, df)
    _escrever_contagem_com_grafico(wb, "Por Cliente", df, "cliente", "Cargas por cliente (destinatário)")
    _escrever_contagem_com_grafico(wb, "Por Remetente", df, "remetente", "Cargas por remetente")
    _escrever_contagem_com_grafico(wb, "Por Motorista", df, "motorista", "Cargas por motorista")
    _escrever_contagem_com_grafico(wb, "Por Cidade", df, "cidade", "Cargas por cidade de entrega")
    _escrever_contagem_com_grafico(
        wb, "Por Status", df, "status", "Cargas por situação", tipo="pizza", limite=10
    )

    caminho_saida = Path(caminho_saida)
    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho_saida)
    logger.info("Relatório visual gerado em: %s", caminho_saida)
    return caminho_saida
