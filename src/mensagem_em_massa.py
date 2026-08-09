"""Gera, a partir de uma planilha de contatos (Destinatário + Telefone),
uma planilha com um link de WhatsApp Web pronto para cada destinatário —
já com a mensagem de posição de cargas preenchida. O operador só abre a
planilha e clica em cada link.
"""
from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src import mensagem
from src.utils import normalizar_texto
from src.whatsapp import montar_link

logger = logging.getLogger("automacao_transportadora")

_FONTE_CABECALHO = Font(color="FFFFFF", bold=True)
_PREENCHIMENTO_CABECALHO = PatternFill("solid", fgColor="1F4E78")
_FONTE_LINK = Font(color="0563C1", underline="single")

_COLUNAS_NOME_CONTATO = ("destinatario", "cliente", "nome", "razao social")
_COLUNAS_TELEFONE_CONTATO = ("telefone", "whatsapp", "celular", "numero", "numero whatsapp")

_COLUNAS_SAIDA = ["destinatario", "telefone", "qtd_cargas", "parte", "link", "observacao"]
_ROTULOS_SAIDA = ["Destinatário", "Telefone", "Qtd. Cargas", "Parte", "Abrir WhatsApp", "Observação"]

# O Excel recusa (ou corrompe) hyperlinks com endereço acima de ~2079
# caracteres. Texto em português (acentos, espaços) quase dobra de
# tamanho ao virar URL (%20, %C3%A7 etc.), e ainda soma o prefixo
# "https://web.whatsapp.com/send?phone=...&text=". Por isso o limite de
# texto por link aqui é bem menor que o usado ao abrir direto no
# navegador (`mensagem.dividir_mensagem`, limite=3500 — sem essa
# restrição do Excel).
_LIMITE_TEXTO_POR_LINK = 900


def _detectar_colunas_contato(df: pd.DataFrame) -> tuple[str, str]:
    colunas_normalizadas = {normalizar_texto(c): c for c in df.columns}

    col_nome = next(
        (colunas_normalizadas[c] for c in _COLUNAS_NOME_CONTATO if c in colunas_normalizadas),
        None,
    )
    col_telefone = next(
        (colunas_normalizadas[c] for c in _COLUNAS_TELEFONE_CONTATO if c in colunas_normalizadas),
        None,
    )

    if not col_nome or not col_telefone:
        raise ValueError(
            "A planilha de contatos precisa ter uma coluna com o nome do "
            "destinatário (ex.: 'Destinatário') e uma com o telefone (ex.: "
            f"'Telefone'). Colunas encontradas: {list(df.columns)}"
        )
    return col_nome, col_telefone


def _cargas_do_destinatario(df_pedidos: pd.DataFrame, nome_destinatario: str) -> pd.DataFrame:
    if "cliente" not in df_pedidos.columns:
        raise ValueError("A planilha de cargas não tem a coluna 'cliente' (Destinatário).")
    alvo = normalizar_texto(nome_destinatario)
    serie = df_pedidos["cliente"].astype(str).map(normalizar_texto)
    return df_pedidos.loc[serie == alvo].copy()


def gerar_planilha_links_whatsapp(
    df_pedidos: pd.DataFrame,
    df_contatos: pd.DataFrame,
    caminho_saida: Path,
) -> Path:
    """Para cada linha de `df_contatos` (destinatário + telefone), busca as
    cargas correspondentes em `df_pedidos` (pela coluna 'cliente'), monta a
    mensagem de posição de cargas (já com o filtro de cargas resolvidas) e
    gera um link de WhatsApp Web pronto. Salva tudo em `caminho_saida`.
    """
    col_nome, col_telefone = _detectar_colunas_contato(df_contatos)

    linhas_saida: list[dict[str, object]] = []

    for _, contato in df_contatos.iterrows():
        nome_destinatario = str(contato.get(col_nome, "") or "").strip()
        telefone = str(contato.get(col_telefone, "") or "").strip()
        if not nome_destinatario:
            continue

        cargas = _cargas_do_destinatario(df_pedidos, nome_destinatario)
        if cargas.empty:
            linhas_saida.append({
                "destinatario": nome_destinatario, "telefone": telefone,
                "qtd_cargas": 0, "parte": "", "link": "",
                "observacao": "Nenhuma carga encontrada para esse destinatário no relatório atual.",
            })
            continue

        cargas_filtradas = mensagem.filtrar_pedidos_para_mensagem(cargas)
        if cargas_filtradas.empty:
            linhas_saida.append({
                "destinatario": nome_destinatario, "telefone": telefone,
                "qtd_cargas": len(cargas), "parte": "", "link": "",
                "observacao": "Todas as cargas já estavam resolvidas/entregues — nenhuma mensagem gerada.",
            })
            continue

        if not telefone:
            linhas_saida.append({
                "destinatario": nome_destinatario, "telefone": "",
                "qtd_cargas": len(cargas_filtradas), "parte": "", "link": "",
                "observacao": "Sem telefone informado na planilha de contatos.",
            })
            continue

        texto = mensagem.montar_mensagem_clientes(cargas_filtradas)
        blocos = mensagem.dividir_mensagem(texto, limite=_LIMITE_TEXTO_POR_LINK)

        for indice, bloco in enumerate(blocos, 1):
            try:
                link = montar_link(telefone, bloco)
                observacao = ""
            except ValueError as erro:
                link = ""
                observacao = f"Telefone inválido: {erro}"

            linhas_saida.append({
                "destinatario": nome_destinatario,
                "telefone": telefone,
                "qtd_cargas": len(cargas_filtradas),
                "parte": f"{indice}/{len(blocos)}" if len(blocos) > 1 else "",
                "link": link,
                "observacao": observacao,
            })

    _salvar_planilha_links(linhas_saida, caminho_saida)
    logger.info(
        "Planilha de links do WhatsApp gerada em: %s (%d linha(s)).",
        caminho_saida, len(linhas_saida),
    )
    return caminho_saida


def _salvar_planilha_links(linhas: list[dict[str, object]], caminho_saida: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Links WhatsApp"

    ws.append(_ROTULOS_SAIDA)
    for celula in ws[1]:
        celula.font = _FONTE_CABECALHO
        celula.fill = _PREENCHIMENTO_CABECALHO

    for linha in linhas:
        ws.append([linha.get(c, "") for c in _COLUNAS_SAIDA])
        if linha.get("link"):
            indice_coluna = _COLUNAS_SAIDA.index("link") + 1
            celula_link = ws.cell(row=ws.max_row, column=indice_coluna)
            celula_link.hyperlink = str(linha["link"])
            celula_link.value = "Abrir WhatsApp"
            celula_link.font = _FONTE_LINK

    larguras = [32, 18, 12, 10, 18, 55]
    for indice, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(indice)].width = largura

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    caminho_saida = Path(caminho_saida)
    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho_saida)
