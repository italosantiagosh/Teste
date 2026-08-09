"""Testes do módulo `relatorio_visual`: garante que a planilha de KPIs e
gráficos é gerada com as abas esperadas, sem quebrar quando faltar alguma
coluna opcional."""

from __future__ import annotations

import sys
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

import pandas as pd
from openpyxl import load_workbook

from src import relatorio_visual


def _df_exemplo() -> pd.DataFrame:
    return pd.DataFrame([
        {
            "pedido": "1001", "cliente": "Cliente A", "remetente": "Remetente A",
            "cidade": "Natal", "status": "Em transito", "motorista": "Fulano",
            "peso_real": "10,5", "volumes": "1", "valor_frete": "100,00",
            "dias_atraso": "0", "duplicado": False,
        },
        {
            "pedido": "1002", "cliente": "Cliente B", "remetente": "Remetente B",
            "cidade": "Mossoro", "status": "Pendente", "motorista": "SEM MANIFESTO",
            "peso_real": "20,0", "volumes": "2", "valor_frete": "200,50",
            "dias_atraso": "3", "duplicado": True,
        },
    ])


def test_gerar_relatorio_visual_cria_abas_esperadas(tmp_path):
    caminho = tmp_path / "analise_visual.xlsx"
    resultado = relatorio_visual.gerar_relatorio_visual(_df_exemplo(), caminho)

    assert resultado == caminho
    assert caminho.exists()

    wb = load_workbook(caminho)
    assert "Resumo" in wb.sheetnames
    assert "Dados" in wb.sheetnames
    assert wb.sheetnames[0] == "Resumo"


def test_gerar_relatorio_visual_dados_mantem_todas_colunas(tmp_path):
    df = _df_exemplo()
    caminho = tmp_path / "analise_visual.xlsx"
    relatorio_visual.gerar_relatorio_visual(df, caminho)

    wb = load_workbook(caminho)
    ws = wb["Dados"]
    cabecalho = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert cabecalho == list(df.columns)
    assert ws.max_row == len(df) + 1


def test_gerar_relatorio_visual_sem_colunas_opcionais_nao_quebra(tmp_path):
    df = pd.DataFrame([{"pedido": "1001", "status": "Em transito"}])
    caminho = tmp_path / "analise_visual.xlsx"
    relatorio_visual.gerar_relatorio_visual(df, caminho)
    assert caminho.exists()
