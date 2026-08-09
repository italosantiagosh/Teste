"""Testes do módulo `mensagem_em_massa`: gera a planilha de links do
WhatsApp para vários destinatários a partir de uma planilha de contatos."""

from __future__ import annotations

import sys
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

import pandas as pd
from openpyxl import load_workbook

from src import mensagem_em_massa


def _df_pedidos() -> pd.DataFrame:
    return pd.DataFrame([
        {
            "pedido": "1001", "cliente": "Comercial Nordeste Ltda",
            "remetente": "Fornecedor A", "pagador": "Pagador X", "cidade": "Natal",
            "nf": "111", "peso_real": "10,0", "volumes": "1", "valor_frete": "50,00",
            "motorista": "Fulano", "status": "Em transito",
        },
        {
            "pedido": "1002", "cliente": "Comercial Nordeste Ltda",
            "remetente": "Fornecedor B", "pagador": "Pagador X", "cidade": "Mossoro",
            "nf": "112", "peso_real": "20,0", "volumes": "2", "valor_frete": "80,00",
            "motorista": "Ciclano", "status": "Entrega realizada ao destinatario.",
        },
        {
            "pedido": "2001", "cliente": "Outra Empresa Ltda",
            "remetente": "Fornecedor C", "pagador": "Pagador Y", "cidade": "Recife",
            "nf": "211", "peso_real": "5,0", "volumes": "1", "valor_frete": "30,00",
            "motorista": "Beltrano", "status": "Em transito",
        },
    ])


def test_gera_link_para_destinatario_com_cargas(tmp_path):
    df_pedidos = _df_pedidos()
    df_contatos = pd.DataFrame([
        {"Destinatário": "Comercial Nordeste Ltda", "Telefone": "5584999341585"},
        {"Destinatário": "Outra Empresa Ltda", "Telefone": "5584999341585"},
    ])

    caminho = tmp_path / "links_whatsapp.xlsx"
    mensagem_em_massa.gerar_planilha_links_whatsapp(df_pedidos, df_contatos, caminho)

    wb = load_workbook(caminho)
    ws = wb.active
    linhas = list(ws.iter_rows(min_row=2, values_only=True))

    por_destinatario = {linha[0]: linha for linha in linhas}
    assert "Comercial Nordeste Ltda" in por_destinatario
    linha_nordeste = por_destinatario["Comercial Nordeste Ltda"]
    # Pedido 1002 já está "entregue" — só a carga 1001 conta.
    assert linha_nordeste[2] == 1
    assert linha_nordeste[4] == "Abrir WhatsApp"

    linha_outra = por_destinatario["Outra Empresa Ltda"]
    assert linha_outra[2] == 1
    assert linha_outra[4] == "Abrir WhatsApp"


def test_destinatario_sem_carga_gera_observacao(tmp_path):
    df_pedidos = _df_pedidos()
    df_contatos = pd.DataFrame([
        {"Destinatário": "Empresa Que Nao Existe", "Telefone": "5584999341585"},
    ])
    caminho = tmp_path / "links_whatsapp.xlsx"
    mensagem_em_massa.gerar_planilha_links_whatsapp(df_pedidos, df_contatos, caminho)

    wb = load_workbook(caminho)
    ws = wb.active
    linha = list(ws.iter_rows(min_row=2, values_only=True))[0]
    assert linha[4] in ("", None)
    assert "Nenhuma carga encontrada" in linha[5]


def test_destinatario_sem_telefone_gera_observacao(tmp_path):
    df_pedidos = _df_pedidos()
    df_contatos = pd.DataFrame([
        {"Destinatário": "Comercial Nordeste Ltda", "Telefone": ""},
    ])
    caminho = tmp_path / "links_whatsapp.xlsx"
    mensagem_em_massa.gerar_planilha_links_whatsapp(df_pedidos, df_contatos, caminho)

    wb = load_workbook(caminho)
    ws = wb.active
    linha = list(ws.iter_rows(min_row=2, values_only=True))[0]
    assert linha[4] in ("", None)
    assert "Sem telefone" in linha[5]


def test_colunas_de_contato_flexiveis(tmp_path):
    df_pedidos = _df_pedidos()
    # Nomes de coluna diferentes, mas reconhecíveis (case/acento não importam).
    df_contatos = pd.DataFrame([
        {"cliente": "Comercial Nordeste Ltda", "whatsapp": "5584999341585"},
    ])
    caminho = tmp_path / "links_whatsapp.xlsx"
    mensagem_em_massa.gerar_planilha_links_whatsapp(df_pedidos, df_contatos, caminho)
    assert caminho.exists()


def test_planilha_contatos_sem_colunas_reconheciveis_da_erro():
    df_pedidos = _df_pedidos()
    df_contatos = pd.DataFrame([{"coluna_qualquer": "valor"}])
    import pytest
    with pytest.raises(ValueError):
        mensagem_em_massa.gerar_planilha_links_whatsapp(df_pedidos, df_contatos, Path("/tmp/nao_usado.xlsx"))


def test_links_ficam_abaixo_do_limite_de_hyperlink_do_excel(tmp_path):
    # O Excel recusa/corrompe hyperlinks com endereço muito longo (na
    # prática, acima de ~2079 caracteres) — com muitas cargas pro mesmo
    # destinatário, a mensagem tem que ser dividida em partes menores
    # especificamente pra esse fluxo (ver _LIMITE_TEXTO_POR_LINK).
    muitas_cargas = pd.DataFrame([
        {
            "pedido": f"10{i:02d}", "cliente": "Cliente Com Muitas Cargas Ltda",
            "remetente": f"Fornecedor Numero {i}", "pagador": "Pagador Unico",
            "cidade": "Sao Goncalo do Amarante", "nf": f"9{i:03d}",
            "peso_real": "123,45", "volumes": "10", "valor_frete": "999,90",
            "motorista": "Motorista Com Nome Bem Grande da Silva Pereira",
            "status": "Chegada na unidade CURRAIS NOVOS em 03/08/26, 10:08h.",
        }
        for i in range(30)
    ])
    df_contatos = pd.DataFrame([
        {"Destinatário": "Cliente Com Muitas Cargas Ltda", "Telefone": "5584999341585"},
    ])

    caminho = tmp_path / "links_whatsapp.xlsx"
    mensagem_em_massa.gerar_planilha_links_whatsapp(muitas_cargas, df_contatos, caminho)

    wb = load_workbook(caminho)
    ws = wb.active
    links = [row[4].hyperlink.target for row in ws.iter_rows(min_row=2) if row[4].hyperlink]
    assert links, "esperava pelo menos um link gerado"
    assert all(len(link) < 2000 for link in links)
